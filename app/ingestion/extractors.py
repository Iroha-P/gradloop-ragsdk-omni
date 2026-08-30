from __future__ import annotations

import re
from pathlib import Path

from .models import ExtractionBlock, ExtractionResult


class ExtractionError(RuntimeError):
    pass


def normalize_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"[ \t]+", " ", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _paragraph_blocks(text: str, *, section: str = "", page: int | None = None) -> list[ExtractionBlock]:
    return [
        ExtractionBlock(text=paragraph.strip(), section=section, page=page)
        for paragraph in re.split(r"\n\s*\n", normalize_text(text))
        if paragraph.strip()
    ]


def extract_markdown(path: Path) -> ExtractionResult:
    text = path.read_text(encoding="utf-8", errors="strict")
    blocks: list[ExtractionBlock] = []
    current_section = path.stem
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.extend(_paragraph_blocks("\n".join(buffer), section=current_section))
            buffer.clear()

    for line in text.splitlines():
        heading = re.match(r"^#{1,6}\s+(.+?)\s*$", line)
        if heading:
            flush()
            current_section = heading.group(1).strip()
        else:
            buffer.append(line)
    flush()
    return ExtractionResult(parser="markdown-stdlib-v1", blocks=blocks)


def extract_text(path: Path) -> ExtractionResult:
    text = path.read_text(encoding="utf-8", errors="strict")
    return ExtractionResult(parser="text-stdlib-v1", blocks=_paragraph_blocks(text, section=path.stem))


def extract_docx(path: Path) -> ExtractionResult:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractionError("DOCX parsing requires python-docx") from exc

    document = Document(str(path))
    blocks: list[ExtractionBlock] = []
    current_section = path.stem
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue
        style_name = getattr(paragraph.style, "name", "") or ""
        if style_name.lower().startswith("heading"):
            current_section = text
            continue
        blocks.append(ExtractionBlock(text=text, section=current_section))
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [normalize_text(cell.text).replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(
                ExtractionBlock(text="\n".join(rows), section=current_section, block_type="table")
            )
    return ExtractionResult(parser="python-docx-v1", blocks=blocks)


def extract_pdf(path: Path, *, max_pages: int = 300, password: str | None = None) -> ExtractionResult:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("PDF parsing requires pypdf") from exc

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        if not password:
            raise ExtractionError("Encrypted PDF requires its configured local password environment variable")
        try:
            decrypted = reader.decrypt(password)
        except Exception as exc:
            raise ExtractionError("Encrypted PDF password could not be applied") from exc
        if decrypted == 0:
            raise ExtractionError("Encrypted PDF password is invalid")
    if len(reader.pages) > max_pages:
        raise ExtractionError("PDF exceeds the configured page limit")

    blocks: list[ExtractionBlock] = []
    warnings: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = normalize_text(page.extract_text() or "")
        except Exception as exc:  # parser errors are isolated to the current file
            raise ExtractionError(f"PDF page {page_number} could not be parsed") from exc
        if not text:
            warnings.append(f"page {page_number} has no extractable text")
            continue
        blocks.extend(_paragraph_blocks(text, section=path.stem, page=page_number))
    return ExtractionResult(parser="pypdf-v1", blocks=blocks, warnings=warnings)


def extract_file(path: Path, *, password: str | None = None) -> ExtractionResult:
    extractor = {
        ".md": extract_markdown,
        ".txt": extract_text,
        ".docx": extract_docx,
        ".pdf": extract_pdf,
    }.get(path.suffix.lower())
    if extractor is None:
        raise ExtractionError(f"Unsupported source type: {path.suffix.lower()}")
    result = extractor(path, password=password) if path.suffix.lower() == ".pdf" else extractor(path)
    if not result.blocks:
        raise ExtractionError("No extractable text was found")
    return result
